import os
import tempfile
import asyncio
import base64
from fastapi import APIRouter, HTTPException, UploadFile, File
from app.core.config import settings

router = APIRouter(prefix="/extract", tags=["extract"])

async def _extract_image_text(file_path: str, mime_type: str) -> str:
    import requests
    def _encode_image():
        with open(file_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode('utf-8')
            
    base64_image = await asyncio.to_thread(_encode_image)
    
    prompt = (
        "You are an expert data extractor. Extract all the text, tables, and important details from this document accurately. "
        "If it is a structured document like an FIR, police report, invoice, or official document, extract the key details clearly (names, dates, incidents, locations). "
        "Preserve table structures using Markdown formatting. "
        "Output ONLY the extracted text and tables in its original language, without any conversational filler."
    )
    
    messages = [
        {
            "role": "user",
            "content": [
                {
                    "type": "text",
                    "text": prompt
                },
                {
                    "type": "image_url",
                    "image_url": {
                        "url": f"data:{mime_type};base64,{base64_image}"
                    }
                }
            ]
        }
    ]
    
    url = "https://openrouter.ai/api/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {settings.OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8000",
        "X-Title": "AI Newsroom Pro"
    }
    payload = {
        "model": settings.OPENROUTER_MODEL,
        "messages": messages,
        "temperature": 0.1
    }
    
    def _request():
        response = requests.post(url, headers=headers, json=payload, timeout=120)
        if not response.ok:
            try:
                err = response.json()
                raise Exception(err.get("error", {}).get("message", response.text))
            except Exception:
                response.raise_for_status()
        data = response.json()
        return data["choices"][0]["message"]["content"].strip()
        
    return await asyncio.to_thread(_request)


@router.post("/", response_model=dict)
async def extract_text_from_file(file: UploadFile = File(...)):
    """Upload a file (PDF, Image, Word, TXT) and extract text"""
    import shutil
    try:
        suffix = os.path.splitext(file.filename)[1]
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            shutil.copyfileobj(file.file, tmp)
            tmp_path = tmp.name

        extracted_text = ""
        file_type = file.content_type.lower()
        suffix_lower = suffix.lower()

        try:
            if "text/plain" in file_type or suffix_lower == '.txt':
                def _read_txt():
                    with open(tmp_path, 'r', encoding='utf-8', errors='ignore') as f:
                        return f.read()
                extracted_text = await asyncio.to_thread(_read_txt)
                
            elif "wordprocessingml.document" in file_type or suffix_lower == '.docx':
                import docx
                def _extract_docx():
                    doc = docx.Document(tmp_path)
                    fullText = []
                    
                    # Extract paragraphs
                    for para in doc.paragraphs:
                        if para.text.strip():
                            fullText.append(para.text.strip())
                            
                    # Extract tables
                    for table in doc.tables:
                        for row in table.rows:
                            row_data = []
                            for cell in row.cells:
                                clean_text = " ".join(cell.text.split())
                                if clean_text:
                                    row_data.append(clean_text)
                            if row_data:
                                fullText.append(" | ".join(row_data))
                                
                    return '\n'.join(fullText)
                extracted_text = await asyncio.to_thread(_extract_docx)
                
            elif "pdf" in file_type or suffix_lower == '.pdf':
                import fitz
                async def _extract_pdf():
                    doc = fitz.open(tmp_path)
                    text_parts = []
                    
                    for page_num in range(len(doc)):
                        page = doc[page_num]
                        page_text = page.get_text()
                        
                        if len(page_text.strip()) > 50:
                            text_parts.append(page_text)
                            try:
                                tables = page.find_tables()
                                if tables:
                                    for table in tables:
                                        df = table.to_pandas()
                                        if df is not None and not df.empty:
                                            text_parts.append(df.to_markdown(index=False))
                            except Exception:
                                pass
                        else:
                            pix = page.get_pixmap(dpi=150)
                            img_path = f"{tmp_path}_page_{page_num}.png"
                            pix.save(img_path)
                            
                            try:
                                img_text = await _extract_image_text(img_path, "image/png")
                                text_parts.append(img_text)
                            except Exception as e:
                                print(f"Vision AI failed for page {page_num}: {e}")
                            finally:
                                if os.path.exists(img_path):
                                    os.remove(img_path)
                                    
                    return "\n\n".join(text_parts)
                extracted_text = await _extract_pdf()
                
            elif "image" in file_type or suffix_lower in ['.png', '.jpg', '.jpeg', '.webp']:
                extracted_text = await _extract_image_text(tmp_path, file_type)
            else:
                raise HTTPException(status_code=400, detail="Unsupported file format. Please upload PDF, Word (.docx), TXT, or an Image.")
        finally:
            if os.path.exists(tmp_path):
                os.remove(tmp_path)

        return {"text": extracted_text}
        
    except HTTPException:
        raise
    except Exception as e:
        print(f"[ERROR] Text extraction failed: {e}")
        raise HTTPException(status_code=500, detail=str(e))
